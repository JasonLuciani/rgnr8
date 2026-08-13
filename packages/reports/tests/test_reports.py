"""Acceptance tests for the reports engine, baseline library, builder, and
renderers."""

from __future__ import annotations

import sqlite3
from datetime import date, datetime
from html import escape
from typing import Any, Callable, cast

import pytest

from rgnr8_forecast import Money

from rgnr8_reports import (
    BASELINE_REPORTS,
    DataContext,
    InMemorySavedReportStore,
    Narrative,
    Report,
    ReportSpec,
    ReportSpecError,
    SqlSavedReportStore,
    Table,
    baseline,
    build_report,
    known_kinds,
    render,
    render_csv,
    render_html,
    spec_to_dict,
    to_dict,
    to_json,
)


# --- baseline coverage -------------------------------------------------------
def test_ten_baselines_present() -> None:
    assert len(BASELINE_REPORTS) == 10
    assert set(BASELINE_REPORTS) == {
        "cash_flow_outlook", "weekly_briefing", "financial_statements", "ar_aging",
        "budget_vs_actual", "runway", "transaction_summary", "usage_billing",
        "exec_board_pack", "reconciliation",
    }
    # every section kind used by a baseline is a registered kind
    for spec in BASELINE_REPORTS.values():
        for section in spec.sections:
            assert section.kind in known_kinds()


@pytest.mark.parametrize("report_id", sorted(BASELINE_REPORTS))
def test_baseline_renders_to_html_and_dict(
    report_id: str,
    full_context: DataContext,
    clock: Callable[[], datetime],
) -> None:
    spec = baseline(report_id)
    report = render(spec, full_context, clock=clock)
    assert isinstance(report, Report)
    assert report.generated_at == clock()
    assert report.sections  # non-empty

    html = render_html(report)
    assert html.startswith("<!doctype html>")
    assert escape(report.title) in html

    d = to_dict(report)
    assert d["title"] == report.title
    # to_json is valid JSON
    json_text = to_json(report)
    assert report.title in json_text


# --- spot-checked numbers ----------------------------------------------------
def test_ar_aging_spot_check(full_context: DataContext, clock: Callable[[], datetime]) -> None:
    report = render(baseline("ar_aging"), full_context, clock=clock)
    html = render_html(report)
    # Total AR is 10,000 + 5,000 + 8,000 = 23,000
    assert "$23,000.00" in html
    # There is an overdue chase list (INV-1 and INV-3 are past due)
    assert "INV-1" in html and "INV-3" in html


def test_financial_statements_spot_check(
    full_context: DataContext, clock: Callable[[], datetime]
) -> None:
    report = render(baseline("financial_statements"), full_context, clock=clock)
    titles = [s.title for s in report.sections]
    assert titles == ["Income Statement", "Balance Sheet", "Cash Flow Statement"]
    html = render_html(report)
    assert "$3,500.00" in html    # net income
    assert "$73,000.00" in html   # total assets
    assert "$50,000.00" in html   # ending cash


def test_transaction_summary_spot_check(
    full_context: DataContext, clock: Callable[[], datetime]
) -> None:
    report = render(baseline("transaction_summary"), full_context, clock=clock)
    html = render_html(report)
    assert "$8,000.00" in html   # money in: 5,000 + 3,000
    assert "$16,000.00" in html  # money out: 4,000 + 12,000


def test_runway_is_finite(full_context: DataContext, clock: Callable[[], datetime]) -> None:
    report = render(baseline("runway"), full_context, clock=clock)
    html = render_html(report)
    assert "months" in html.lower()


def test_reconciliation_spot_check(
    full_context: DataContext, clock: Callable[[], datetime]
) -> None:
    report = render(baseline("reconciliation"), full_context, clock=clock)
    html = render_html(report)
    # ledger 50,000 vs bank 49,500 → 500 gap → MINOR
    assert "MINOR" in html
    assert "$500.00" in html


def test_usage_billing_invoice_preview(
    full_context: DataContext, clock: Callable[[], datetime]
) -> None:
    report = render(baseline("usage_billing"), full_context, clock=clock)
    html = render_html(report)
    # Assisted plan: base 499 + 80 overage analyst minutes @ 2.50 = 200 → 699 total
    assert "$699.00" in html
    assert "200" in html  # analyst minutes


# --- CSV ---------------------------------------------------------------------
def test_csv_flattens_tables(full_context: DataContext, clock: Callable[[], datetime]) -> None:
    report = render(baseline("ar_aging"), full_context, clock=clock)
    csv_text = render_csv(report)
    lines = csv_text.splitlines()
    # section header + column header appear; a bucket row is present
    assert "Aging Summary" in csv_text
    assert "Bucket,Open amount" in csv_text
    assert any(line.startswith("Current,") for line in lines)


def test_csv_multiple_tables_separated(
    full_context: DataContext, clock: Callable[[], datetime]
) -> None:
    report = render(baseline("financial_statements"), full_context, clock=clock)
    csv_text = render_csv(report)
    # three tables, each prefixed by its section header, blank-line separated
    assert "Income Statement" in csv_text
    assert "Balance Sheet" in csv_text
    assert "Cash Flow Statement" in csv_text
    assert "\r\n\r\n" in csv_text  # blank-line separator between tables


# --- self-contained HTML -----------------------------------------------------
def test_html_is_self_contained(
    full_context: DataContext, clock: Callable[[], datetime]
) -> None:
    for report_id in BASELINE_REPORTS:
        report = render(baseline(report_id), full_context, clock=clock)
        html = render_html(report)
        assert "http://" not in html
        assert "https://" not in html
        assert "<style>" in html  # CSS inlined


# --- graceful degradation ----------------------------------------------------
def test_missing_data_degrades_to_not_available(clock: Callable[[], datetime]) -> None:
    empty = DataContext(period="Empty", as_of=date(2026, 8, 1))
    report = render(baseline("financial_statements"), empty, clock=clock)
    # every section falls back to a single "Not available" narrative
    for section in report.sections:
        assert len(section.blocks) == 1
        block = section.blocks[0]
        assert isinstance(block, Narrative)
        assert "Not available" in block.text
    # and it still renders to HTML without error
    assert "Not available" in render_html(report)


# --- custom builder ----------------------------------------------------------
def test_build_report_from_section_list(
    full_context: DataContext, clock: Callable[[], datetime]
) -> None:
    spec_json = [
        {"kind": "narrative", "title": "Intro", "params": {"text": "Hello owner."}},
        {"kind": "ar_aging", "title": "Receivables"},
        {"kind": "kpi_row", "title": "Highlights",
         "params": {"items": [["Cash", "$50,000.00", False], ["Risk", "-$500.00", True]]}},
    ]
    spec = build_report(spec_json)
    assert isinstance(spec, ReportSpec)
    assert [s.kind for s in spec.sections] == ["narrative", "ar_aging", "kpi_row"]

    report = render(spec, full_context, clock=clock)
    html = render_html(report)
    assert "Hello owner." in html
    assert "$23,000.00" in html  # ar_aging rendered against the context


def test_build_report_rejects_unknown_kind() -> None:
    with pytest.raises(ReportSpecError):
        build_report([{"kind": "does_not_exist", "title": "x"}])


def test_build_report_accepts_full_object() -> None:
    spec = build_report({
        "id": "my_report",
        "title": "My Report",
        "description": "custom",
        "sections": [{"kind": "runway"}],
    })
    assert spec.id == "my_report"
    assert spec.title == "My Report"


# --- saved report stores -----------------------------------------------------
def _sample_custom_spec() -> ReportSpec:
    return build_report({
        "id": "custom_pack",
        "title": "Custom Pack",
        "sections": [
            {"kind": "cash_outlook", "title": "Cash"},
            {"kind": "narrative", "params": {"text": "note"}},
        ],
    })


def test_inmemory_store_roundtrip() -> None:
    store = InMemorySavedReportStore()
    spec = _sample_custom_spec()
    store.save("tenant-a", spec)
    loaded = store.get("tenant-a", "custom_pack")
    assert loaded is not None
    assert spec_to_dict(loaded) == spec_to_dict(spec)
    assert store.get("tenant-b", "custom_pack") is None
    assert [s.id for s in store.list_for_tenant("tenant-a")] == ["custom_pack"]


def test_sql_store_roundtrip_sqlite(
    full_context: DataContext, clock: Callable[[], datetime]
) -> None:
    conn = sqlite3.connect(":memory:")
    try:
        store = SqlSavedReportStore(cast(Any, conn), placeholder="?")
        store.create_schema()
        spec = _sample_custom_spec()
        store.save("tenant-a", spec)

        loaded = store.get("tenant-a", "custom_pack")
        assert loaded is not None
        assert spec_to_dict(loaded) == spec_to_dict(spec)

        # tenant isolation
        assert store.get("tenant-b", "custom_pack") is None
        assert [s.id for s in store.list_for_tenant("tenant-a")] == ["custom_pack"]

        # a reloaded spec still renders
        report = render(loaded, full_context, clock=clock)
        assert render_html(report).startswith("<!doctype html>")
    finally:
        conn.close()


# --- block model -------------------------------------------------------------
def test_to_dict_block_kinds(full_context: DataContext, clock: Callable[[], datetime]) -> None:
    report = render(baseline("cash_flow_outlook"), full_context, clock=clock)
    d = to_dict(report)
    sections = cast("list[dict[str, Any]]", d["sections"])
    kinds: set[str] = set()
    for section in sections:
        for block in cast("list[dict[str, Any]]", section["blocks"]):
            kinds.add(cast(str, block["kind"]))
    # cash_outlook produces a narrative, a kpi row, a table, and a chart
    assert {"narrative", "kpi_row", "table", "chart"} <= kinds


# --- PDF export --------------------------------------------------------------
def test_pdf_export_is_valid_and_deterministic(
    full_context: DataContext, clock: Callable[[], datetime]
) -> None:
    from rgnr8_reports import render_pdf

    report = render(baseline("exec_board_pack"), full_context, clock=clock)
    a = render_pdf(report)
    b = render_pdf(report)
    assert a.startswith(b"%PDF-")
    assert a.endswith(b"%%EOF\n") or a.rstrip().endswith(b"%%EOF")
    # Byte-deterministic: reportlab invariant mode pins timestamp + doc id.
    assert a == b
    assert len(a) > 1000


@pytest.mark.parametrize("report_id", sorted(BASELINE_REPORTS))
def test_every_baseline_renders_to_pdf(
    report_id: str, full_context: DataContext, clock: Callable[[], datetime]
) -> None:
    from rgnr8_reports import render_pdf

    report = render(baseline(report_id), full_context, clock=clock)
    data = render_pdf(report)
    assert data.startswith(b"%PDF-")


def test_pdf_of_empty_context_still_renders(clock: Callable[[], datetime]) -> None:
    from rgnr8_reports import render_pdf

    empty = DataContext(period="Empty", as_of=date(2026, 8, 1))
    report = render(baseline("financial_statements"), empty, clock=clock)
    assert render_pdf(report).startswith(b"%PDF-")


# --- XLSX export -------------------------------------------------------------
def test_xlsx_export_has_cover_and_table_sheets(
    full_context: DataContext, clock: Callable[[], datetime]
) -> None:
    import io

    import openpyxl

    from rgnr8_reports import render_xlsx

    report = render(baseline("exec_board_pack"), full_context, clock=clock)
    data = render_xlsx(report)
    assert data[:2] == b"PK"  # zip container
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames[0] == "Overview"
    # every data table becomes its own worksheet
    table_sections = [
        s.title
        for s in report.sections
        for b in s.blocks
        if b.__class__.__name__ == "Table"
    ]
    for title in table_sections:
        assert any(title[:31] == name or title[:28] in name for name in wb.sheetnames)


def test_xlsx_writes_amounts_as_numbers(
    full_context: DataContext, clock: Callable[[], datetime]
) -> None:
    import io

    import openpyxl

    from rgnr8_reports import render_xlsx

    report = render(baseline("financial_statements"), full_context, clock=clock)
    wb = openpyxl.load_workbook(io.BytesIO(render_xlsx(report)))
    # find at least one numeric amount cell in a table sheet
    found_number = False
    for name in wb.sheetnames[1:]:
        for row in wb[name].iter_rows():
            for cell in row:
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    found_number = True
    assert found_number


def test_xlsx_as_number_parser() -> None:
    from rgnr8_reports.xlsx import _as_number

    assert _as_number("$10,000.00") == 10000.0
    assert _as_number("-$4,000.00") == -4000.0
    assert _as_number("EUR 1,250.00") == 1250.0
    assert _as_number("4.2 mo") is None
    assert _as_number("W1") is None
    assert _as_number("") is None


@pytest.mark.parametrize("report_id", sorted(BASELINE_REPORTS))
def test_every_baseline_renders_to_xlsx(
    report_id: str, full_context: DataContext, clock: Callable[[], datetime]
) -> None:
    from rgnr8_reports import render_xlsx

    report = render(baseline(report_id), full_context, clock=clock)
    assert render_xlsx(report)[:2] == b"PK"


# --- scheduled reports -------------------------------------------------------
def _sched(fmt: str = "pdf", report_id: str = "exec_board_pack"):
    from zoneinfo import ZoneInfo  # noqa: F401

    from rgnr8_briefing import Schedule

    from rgnr8_reports import ReportSchedule

    return ReportSchedule(
        tenant_id="t1", report_id=report_id, recipient="owner@acme.com",
        schedule=Schedule(weekday=0, hour=8, minute=0, timezone="America/Denver"), fmt=fmt,
    )


def _tz_now(day: int, hour: int = 15):
    from zoneinfo import ZoneInfo

    return datetime(2026, 8, day, hour, 0, 0, tzinfo=ZoneInfo("UTC"))


def test_schedule_fires_once_per_cadence_with_catch_up(full_context: DataContext) -> None:
    from rgnr8_reports import RecordingReportSink, run_due_reports

    scheds = [_sched()]
    sink = RecordingReportSink()
    kw = dict(spec_for=lambda s: baseline(s.report_id), context_for=lambda s: full_context, sink=sink)

    # Monday after 8am Denver -> due
    o1 = run_due_reports(scheds, _tz_now(10), **kw)  # 2026-08-10 is a Monday
    assert [o.fired for o in o1] == [True]
    assert len(sink.sent) == 1
    # same week, later -> not due (idempotent)
    o2 = run_due_reports(scheds, _tz_now(10, 18), **kw)
    assert o2[0].fired is False and o2[0].skipped_reason == "not_due"
    assert len(sink.sent) == 1
    # next week -> fires again exactly once
    o3 = run_due_reports(scheds, _tz_now(17), **kw)
    assert o3[0].fired is True
    assert len(sink.sent) == 2


def test_schedule_delivers_binary_pdf_and_text_csv(full_context: DataContext) -> None:
    from rgnr8_reports import RecordingReportSink, run_due_reports

    for fmt, head in (("pdf", b"%PDF-"), ("xlsx", b"PK")):
        sink = RecordingReportSink()
        run_due_reports(
            [_sched(fmt=fmt)], _tz_now(10),
            spec_for=lambda s: baseline(s.report_id),
            context_for=lambda s: full_context, sink=sink,
        )
        d = sink.sent[0]
        assert d.is_binary and isinstance(d.content, bytes) and d.content.startswith(head)
        assert d.filename == f"exec_board_pack.{fmt}"

    sink = RecordingReportSink()
    run_due_reports(
        [_sched(fmt="csv")], _tz_now(10),
        spec_for=lambda s: baseline(s.report_id),
        context_for=lambda s: full_context, sink=sink,
    )
    assert not sink.sent[0].is_binary and isinstance(sink.sent[0].content, str)


def test_schedule_paused_and_not_ready_skip(full_context: DataContext) -> None:
    from rgnr8_reports import RecordingReportSink, run_due_reports

    paused = _sched()
    paused.active = False
    sink = RecordingReportSink()
    out = run_due_reports(
        [paused], _tz_now(10),
        spec_for=lambda s: baseline(s.report_id),
        context_for=lambda s: full_context, sink=sink,
    )
    assert out[0].skipped_reason == "paused" and not sink.sent

    # spec/context not ready -> skipped "not_ready", nothing sent
    sink2 = RecordingReportSink()
    out2 = run_due_reports(
        [_sched()], _tz_now(10),
        spec_for=lambda s: None, context_for=lambda s: full_context, sink=sink2,
    )
    assert out2[0].skipped_reason == "not_ready" and not sink2.sent


def test_schedule_rejects_unknown_format() -> None:
    import pytest as _pytest

    with _pytest.raises(ValueError):
        _sched(fmt="docx")


def test_render_in_format_dispatch(full_context: DataContext, clock: Callable[[], datetime]) -> None:
    from rgnr8_reports import render_in_format

    report = render(baseline("exec_board_pack"), full_context, clock=clock)
    assert isinstance(render_in_format(report, "html"), str)
    assert isinstance(render_in_format(report, "json"), str)
    assert isinstance(render_in_format(report, "csv"), str)
    assert isinstance(render_in_format(report, "pdf"), bytes)
    assert isinstance(render_in_format(report, "xlsx"), bytes)
    with __import__("pytest").raises(ValueError):
        render_in_format(report, "nope")


def test_schedule_store_roundtrip_inmemory() -> None:
    from rgnr8_reports import InMemoryReportScheduleStore

    store = InMemoryReportScheduleStore()
    store.save(_sched(report_id="a"))
    store.save(_sched(report_id="b"))
    store.save(ReportScheduleForOtherTenant())
    assert len(store.list_for_tenant("t1")) == 2
    assert len(store.list_all()) == 3
    assert store.get("t1", "a") is not None
    store.delete("t1", "a")
    assert store.get("t1", "a") is None
    assert len(store.list_for_tenant("t1")) == 1


def ReportScheduleForOtherTenant():  # noqa: N802 - tiny helper
    from rgnr8_briefing import Schedule

    from rgnr8_reports import ReportSchedule

    return ReportSchedule(
        tenant_id="t2", report_id="a", recipient="o@t2.com",
        schedule=Schedule(), fmt="pdf",
    )


def test_schedule_store_roundtrip_sqlite() -> None:
    from rgnr8_reports import SqlReportScheduleStore

    conn = sqlite3.connect(":memory:")
    try:
        store = SqlReportScheduleStore(conn)
        store.create_schema()
        s = _sched()
        s.last_sent = datetime(2026, 8, 10, 8, 0, 0)
        store.save(s)
        loaded = store.get("t1", "exec_board_pack")
        assert loaded is not None
        assert loaded.fmt == "pdf"
        assert loaded.last_sent == datetime(2026, 8, 10, 8, 0, 0)
        assert loaded.schedule.timezone == "America/Denver"
        assert len(store.list_all()) == 1
        store.delete("t1", "exec_board_pack")
        assert store.get("t1", "exec_board_pack") is None
    finally:
        conn.close()
