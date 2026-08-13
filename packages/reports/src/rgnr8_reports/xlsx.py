"""XLSX renderer — turn a rendered :class:`Report` into a branded workbook.

Owners and their accountants live in spreadsheets, so a report needs an Excel
export they can pivot, chart, and drop into their own models. This renders every
report to a single ``.xlsx`` workbook: a cover sheet with the title, period, and
each section's KPI figures, then one worksheet per data :class:`Table` (named for
its section) so a multi-section report opens as a tidy set of tabs.

Built with :mod:`openpyxl` from the structured :class:`Report` model. Formatting
follows the RGNR8 brand (forest header fills, ivory bands, the risk token for
negatives); numeric columns are right-aligned and, where a cell parses as a
currency amount, written as a real number with an accounting number format so the
sheet totals and charts correctly rather than carrying text. Values that are not
numeric (labels, dates, status) are written as text verbatim.

:class:`~rgnr8_reports.model.Chart` blocks are inline SVG for the web surface and
carry no figures a table doesn't, so they are omitted from the workbook.
"""

from __future__ import annotations

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from rgnr8_forecast.brand import FOREST, IVORY, INK, LINE, MUTED, RISK

from .model import KpiRow, Report, Table

# --- brand styles (openpyxl wants ARGB hex, no leading '#') -------------------
def _argb(hex_color: str) -> str:
    return "FF" + hex_color.lstrip("#").upper()


_FONT = "Arial"
_HEADER_FILL = PatternFill("solid", fgColor=_argb(FOREST))
_BAND_FILL = PatternFill("solid", fgColor=_argb(IVORY))
_HEADER_FONT = Font(name=_FONT, bold=True, size=9, color=_argb(IVORY))
_TITLE_FONT = Font(name=_FONT, bold=True, size=16, color=_argb(INK))
_LABEL_FONT = Font(name=_FONT, bold=True, size=8, color=_argb(MUTED))
_BODY_FONT = Font(name=_FONT, size=10, color=_argb(INK))
_NEG_FONT = Font(name=_FONT, size=10, color=_argb(RISK))
_META_FONT = Font(name=_FONT, size=9, color=_argb(MUTED))
_THIN = Side(style="thin", color=_argb(LINE))
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_RIGHT = Alignment(horizontal="right")
_LEFT = Alignment(horizontal="left")

# Accounting format: negatives in parentheses, zero as a dash — matches finance.
_MONEY_FMT = '#,##0.00;(#,##0.00);"-"'
_MONEY_RE = re.compile(r"^-?[\$€£]?\s*[\d,]+(?:\.\d+)?$|^-?[A-Z]{3}\s*[\d,]+(?:\.\d+)?$")


def _is_negish(value: str) -> bool:
    return value.strip().startswith("-") or value.strip().startswith("(")


def _as_number(value: str) -> float | None:
    """If ``value`` reads as a currency/number cell, its numeric value; else None."""
    v = value.strip()
    if not _MONEY_RE.match(v):
        return None
    neg = v.startswith("-") or (v.startswith("(") and v.endswith(")"))
    cleaned = re.sub(r"[^\d.]", "", v)
    if cleaned in ("", "."):
        return None
    try:
        num = float(cleaned)
    except ValueError:
        return None
    return -num if neg else num


def _safe_sheet_title(title: str, used: set[str]) -> str:
    """A valid, unique worksheet name (<=31 chars, no ``[]:*?/\\``)."""
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", title).strip() or "Sheet"
    cleaned = cleaned[:31]
    candidate = cleaned
    n = 2
    while candidate.lower() in used:
        suffix = f" {n}"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def _write_table(ws: Worksheet, block: Table, start_row: int) -> int:
    """Write a table block starting at ``start_row``; return the next free row."""
    row = start_row
    # header
    for c, col in enumerate(block.columns, start=1):
        cell = ws.cell(row=row, column=c, value=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.border = _BORDER
        cell.alignment = _RIGHT if c > 1 else _LEFT
    row += 1
    for r_i, data_row in enumerate(block.rows):
        for c, raw in enumerate(data_row, start=1):
            num = _as_number(raw) if c > 1 else None
            if num is not None:
                cell = ws.cell(row=row, column=c, value=num)
                cell.number_format = _MONEY_FMT
                cell.alignment = _RIGHT
            else:
                cell = ws.cell(row=row, column=c, value=raw)
                cell.alignment = _RIGHT if c > 1 else _LEFT
            cell.font = _NEG_FONT if _is_negish(raw) else _BODY_FONT
            cell.border = _BORDER
            if r_i % 2 == 1:
                cell.fill = _BAND_FILL
        row += 1
    return row + 1


def _autosize(ws: Worksheet) -> None:
    """Widen columns to fit their longest cell (deterministic, content-based)."""
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            length = len(str(cell.value))
            widths[cell.column] = max(widths.get(cell.column, 0), length)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 3, 10), 48)


def render_xlsx(report: Report) -> bytes:
    """Render ``report`` as a branded ``.xlsx`` workbook (cover + one tab per table)."""
    wb = Workbook()
    used: set[str] = set()

    cover = wb.active
    assert cover is not None
    cover.title = _safe_sheet_title("Overview", used)
    cover.sheet_view.showGridLines = False

    cover["A1"] = report.title
    cover["A1"].font = _TITLE_FONT
    meta = f"Generated {report.generated_at.isoformat()}"
    if report.period:
        meta += f"  ·  {report.period}"
    cover["A2"] = meta
    cover["A2"].font = _META_FONT
    row = 4

    # KPI figures from every section, listed on the cover.
    for section in report.sections:
        kpis = [b for b in section.blocks if isinstance(b, KpiRow)]
        if not kpis:
            continue
        cover.cell(row=row, column=1, value=section.title).font = Font(
            name=_FONT, bold=True, size=11, color=_argb(FOREST)
        )
        row += 1
        for kpi in kpis:
            for label, value, is_negative in kpi.items:
                lc = cover.cell(row=row, column=1, value=label)
                lc.font = _LABEL_FONT
                vnum = _as_number(value)
                if vnum is not None:
                    vc = cover.cell(row=row, column=2, value=vnum)
                    vc.number_format = _MONEY_FMT
                else:
                    vc = cover.cell(row=row, column=2, value=value)
                vc.font = _NEG_FONT if is_negative else _BODY_FONT
                vc.alignment = _RIGHT
                row += 1
        row += 1
    _autosize(cover)

    # One worksheet per data table.
    for section in report.sections:
        for block in section.blocks:
            if not isinstance(block, Table):
                continue
            ws = wb.create_sheet(_safe_sheet_title(section.title, used))
            ws.sheet_view.showGridLines = False
            ws.cell(row=1, column=1, value=section.title).font = Font(
                name=_FONT, bold=True, size=13, color=_argb(INK)
            )
            _write_table(ws, block, 3)
            _autosize(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
